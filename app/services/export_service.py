import json
import csv
import io
from pathlib import Path
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.record import VoterRecord
from app.config import EXPORT_DIR

def export_to_json(records: List[VoterRecord], filter_type: str = "all") -> str:
    """
    Exports voter records to JSON string.
    filter_type: 'all', 'verified_only', 'needs_review_only'
    """
    filtered = filter_records(records, filter_type)
    data = []
    for r in filtered:
        item = {
            "serial_number": r.serial_number,
            "part_number": r.part_number,
            "voter_id": r.voter_id,
            "elector_name": r.elector_name,
            "relation_type": r.relation_type,
            "relation_name": r.relation_name,
            "house_number": r.house_number,
            "age": r.age,
            "gender": r.gender,
            "gender_original": r.gender_original,
            "photo_available": r.photo_available,
            "source_page": r.source_page,
            "card_index": r.card_index,
            "verification_status": r.verification_status,
            "needs_review": r.needs_review,
            "confidence": r.confidence,
            "review_reasons": r.review_reasons
        }
        if r.source_bbox:
            item["source_bbox"] = {
                "x": r.source_bbox.x,
                "y": r.source_bbox.y,
                "width": r.source_bbox.width,
                "height": r.source_bbox.height
            }
        data.append(item)
        
    return json.dumps(data, ensure_ascii=False, indent=2)

def export_to_xlsx(records: List[VoterRecord], filter_type: str = "all") -> bytes:
    """
    Generates an Excel workbook with frozen headers, styled header row,
    autofilter, and auto-adjusted column widths.
    """
    filtered = filter_records(records, filter_type)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Voter Records"
    
    headers = [
        "Serial Number",
        "Part Number",
        "Voter ID",
        "Elector Name",
        "Relation Type",
        "Relation Name",
        "House Number",
        "Age",
        "Gender",
        "Gender (Original)",
        "Photo Available",
        "Source Page",
        "Card Index",
        "Verification Status",
        "Needs Review",
        "Review Reasons"
    ]
    ws.append(headers)
    
    # Style Header
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = thin_border
        
    # Add Data
    for r in filtered:
        row_values = [
            r.serial_number if r.serial_number is not None else "",
            r.part_number if r.part_number is not None else "",
            r.voter_id or "",
            r.elector_name or "",
            r.relation_type or "",
            r.relation_name or "",
            r.house_number or "",
            r.age if r.age is not None else "",
            r.gender or "",
            r.gender_original or "",
            "Yes" if r.photo_available else "No",
            r.source_page,
            r.card_index,
            r.verification_status,
            "Yes" if r.needs_review else "No",
            "; ".join(r.review_reasons) if r.review_reasons else ""
        ]
        ws.append(row_values)
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Freeze top row and enable autofilter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def export_to_csv(records: List[VoterRecord], filter_type: str = "all") -> bytes:
    """
    Generates CSV with utf-8-sig encoding for seamless Hindi viewing in Excel.
    """
    filtered = filter_records(records, filter_type)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    
    headers = [
        "Serial Number",
        "Part Number",
        "Voter ID",
        "Elector Name",
        "Relation Type",
        "Relation Name",
        "House Number",
        "Age",
        "Gender",
        "Gender Original",
        "Photo Available",
        "Source Page",
        "Card Index",
        "Verification Status",
        "Needs Review"
    ]
    writer.writerow(headers)
    
    for r in filtered:
        writer.writerow([
            r.serial_number if r.serial_number is not None else "",
            r.part_number if r.part_number is not None else "",
            r.voter_id or "",
            r.elector_name or "",
            r.relation_type or "",
            r.relation_name or "",
            r.house_number or "",
            r.age if r.age is not None else "",
            r.gender or "",
            r.gender_original or "",
            "true" if r.photo_available else "false",
            r.source_page,
            r.card_index,
            r.verification_status,
            "true" if r.needs_review else "false"
        ])
        
    return buffer.getvalue().encode("utf-8-sig")

def export_to_aligned_txt(records: List[VoterRecord], filter_type: str = "all") -> str:
    """
    Generates tabular aligned text with dynamically calculated column widths.
    """
    filtered = filter_records(records, filter_type)
    if not filtered:
        return "No records extracted."
        
    cols = ["Serial", "Voter ID", "Elector Name", "Relation", "Relation Name", "House", "Age", "Gender", "Page", "Status"]
    
    rows = []
    for r in filtered:
        rel_summary = f"{r.relation_type or ''}".capitalize()
        rows.append([
            str(r.serial_number or "-"),
            r.voter_id or "-",
            r.elector_name or "-",
            rel_summary,
            r.relation_name or "-",
            r.house_number or "-",
            str(r.age or "-"),
            r.gender_original or r.gender or "-",
            str(r.source_page),
            r.verification_status
        ])
        
    # Calculate max column widths
    widths = [len(c) for c in cols]
    for row in rows:
        for i, val in enumerate(row):
            widths[i] = max(widths[i], len(val))
            
    header_line = " | ".join(c.ljust(widths[i]) for i, c in enumerate(cols))
    sep_line = "-+-".join("-" * widths[i] for i in range(len(cols)))
    
    output_lines = [header_line, sep_line]
    for row in rows:
        output_lines.append(" | ".join(row[i].ljust(widths[i]) for i in range(len(cols))))
        
    return "\n".join(output_lines)

def filter_records(records: List[VoterRecord], filter_type: str) -> List[VoterRecord]:
    if filter_type == "verified_only":
        return [r for r in records if r.verification_status == "verified"]
    elif filter_type == "needs_review_only":
        return [r for r in records if r.needs_review]
    return records
